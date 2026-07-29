"""Generate NL_Blinkit_CategoryAdoption.pptx from design/Blinkit Category Nudge Case Study.dc.html.

An editable, native-text PowerPoint counterpart to the HTML deck (which is the source of truth —
see design/CLAUDE.md for the hard rules on content/figures). This script mirrors that HTML's
copy, figures, and visual system slide-for-slide so PowerPoint users get real text boxes and
tables instead of a flattened image export. Re-run after any content change to the HTML deck.

Run:  ./venv/bin/python deck/build_deck.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "design" / "img"

# ---- visual system (design/CLAUDE.md) ----
INK = RGBColor(0x16, 0x13, 0x0A)
PAGE = RGBColor(0xF4, 0xF5, 0xF3)
YELLOW = RGBColor(0xF8, 0xCD, 0x1B)
AMBER_TEXT = RGBColor(0x7A, 0x61, 0x00)
GREEN_TEXT = RGBColor(0x14, 0x66, 0x34)
GREEN_ACCENT = RGBColor(0x31, 0x86, 0x16)
BODY = RGBColor(0x3A, 0x36, 0x28)
MUTED = RGBColor(0x5A, 0x55, 0x45)
FOOTNOTE = RGBColor(0x6B, 0x68, 0x57)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xE7, 0xE8, 0xE2)
LIGHT_GREEN_BG = RGBColor(0xED, 0xF5, 0xE9)
LIGHT_GREEN_BORDER = RGBColor(0xCF, 0xE3, 0xC6)
LIGHT_AMBER_BG = RGBColor(0xFF, 0xF9, 0xDF)
LIGHT_AMBER_BORDER = RGBColor(0xF0, 0xDF, 0xA0)
DARK_TEXT_MUTED = RGBColor(0xC9, 0xC7, 0xBC)
DARK_TEXT_MUTED2 = RGBColor(0xDA, 0xD8, 0xCE)

FONT = "Arial"  # Plus Jakarta Sans isn't guaranteed on every machine; Arial is the safe substitute
MIN = 14  # never render below this (Google Slides / PPT floor)

DISCOVERY_URL = "https://huggingface.co/spaces/Abhishek292000/blinkit-discovery-engine"
MVP_URL = "https://huggingface.co/spaces/Abhishek292000/blinkit-category-nudge-agent"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAGE
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf


def run(p, text, size, color=INK, bold=False, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(max(size, MIN))  # brief's hard floor: never render below 14pt
    r.font.name = FONT
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    return r


def panel(s, x, y, w, h, fill=None, border=CARD_BORDER, line_w=0.75):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if border is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = border
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def eyebrow(s, text, x=0.55, y=0.35, color=AMBER_TEXT, with_logo=True, size=13):
    logo_w = 0.85
    lx = x
    if with_logo and (IMG / "blinkit-logo.png").exists():
        s.shapes.add_picture(str(IMG / "blinkit-logo.png"), Inches(lx), Inches(y), height=Inches(0.3))
        lx += logo_w
    _, tf = box(s, lx, y - 0.03, 10.5, 0.4, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, text.upper(), size, color, bold=True)


def h2(s, text, y=0.78, size=26, w=11.6, color=INK):
    _, tf = box(s, 0.55, y, w, 1.35)
    p = tf.paragraphs[0]
    run(p, text, size, color, bold=True)


def bullet_para(tf, text, size=14, color=BODY, bold=False, first=False, space_after=6, bullet="•  "):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    run(p, bullet + text, size, color, bold=bold)
    return p


def footnote(s, text, x=0.55, y=7.05, w=12.2, size=11, color=FOOTNOTE):
    _, tf = box(s, x, y, w, 0.4)
    p = tf.paragraphs[0]
    run(p, text, size, color)


def simple_table(s, x, y, w, h, headers, rows, col_widths, header_fill=INK, header_color=WHITE,
                  body_size=12, header_size=13, row_fill=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gtable = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = gtable.table
    total = sum(col_widths)
    for i, cw in enumerate(col_widths):
        table.columns[i].width = Inches(w * cw / total)
    for j, htext in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run(p, htext, header_size, header_color, bold=True)
    for i, row in enumerate(rows):
        fill = None
        if row_fill and i in row_fill:
            fill = row_fill[i]
        for j, (text, color, bold) in enumerate(row):
            cell = table.cell(i + 1, j)
            if fill is not None:
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run(p, text, body_size, color, bold=bold)
    return table


def picture_fit(s, path, x, y, max_w, max_h):
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    if max_w / max_h > ar:
        h = max_h
        w = h * ar
    else:
        w = max_w
        h = w / ar
    px = x + (max_w - w) / 2
    py = y + (max_h - h) / 2
    return s.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(w), height=Inches(h))


# ========================================================================
# SLIDE 1 — Title
# ========================================================================
s = slide()
eyebrow(s, "Product case study · Blinkit · quick commerce", y=0.5, with_logo=True, size=13)
_, tf = box(s, 0.55, 0.95, 12.2, 1.15)
p = tf.paragraphs[0]
run(p, "Why do Blinkit's most loyal users never widen their basket?", 32, INK, bold=True)
_, tf2 = box(s, 0.55, 2.05, 12.2, 0.8)
p2 = tf2.paragraphs[0]
run(p2, "Four parts: an AI discovery engine over public reviews, primary user research, a "
        "problem definition, and a deployed AI-native MVP — the Category Nudge Agent.", 15, BODY)

# left dark panel — customer base stats
panel(s, 0.55, 2.95, 6.15, 4.15, fill=INK, border=None)
_, tf = box(s, 0.85, 3.15, 5.6, 0.4)
run(tf.paragraphs[0], "THE CUSTOMER BASE WE ARE DESIGNING INTO", 12, YELLOW, bold=True)
stats = [
    ("31.8M", "avg. monthly transacting customers", "16.9M a year earlier", YELLOW),
    ("331M", "orders in the quarter", "≈ 10.4 orders per customer", WHITE),
    ("2,443", "dark stores live", "+200 in the quarter", WHITE),
    ("86%", "YoY net order value growth", "₹17,132 cr NOV", WHITE),
]
col_w = 1.42
for i, (val, label, sub, color) in enumerate(stats):
    x = 0.85 + i * col_w
    _, tf = box(s, x, 3.65, col_w - 0.1, 1.7)
    p = tf.paragraphs[0]
    run(p, val, 26, color, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(3)
    run(p2, label, 11, WHITE, bold=True)
    p3 = tf.add_paragraph(); p3.space_before = Pt(2)
    run(p3, sub, 10, DARK_TEXT_MUTED)
_, tf = box(s, 0.85, 5.75, 5.6, 0.85)
run(tf.paragraphs[0], "Blinkit segment, Eternal Q1 FY27 (quarter ended 30 Jun 2026) shareholders' "
                      "letter and earnings call. Growth came from more customers and more orders, "
                      "not higher basket value — average order value dipped slightly.", 10, DARK_TEXT_MUTED)

# right white panel — discovery on Blinkit today
panel(s, 6.9, 2.95, 5.9, 4.15, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 7.15, 3.12, 5.4, 0.35)
run(tf.paragraphs[0], "DISCOVERY ON BLINKIT TODAY", 12, AMBER_TEXT, bold=True)
_, tf = box(s, 7.15, 3.5, 5.4, 1.95)
items = [
    ("Intent-led search and category aisles.", "AI-assisted search ranking is the primary way a basket gets filled — the customer has to ask first."),
    ("Home rails tuned to place and time.", "What is popular nearby, time-of-day sets, past purchases, and offer rails on things the customer already looked at."),
    ("Complement prompts at cart.", "Add-on suggestions next to what is already in the basket — adjacent to intent, rarely a new category."),
]
first = True
for i, (bold_t, rest) in enumerate(items):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(5)
    run(p, f"{i+1}.  ", 12, GREEN_TEXT, bold=True)
    run(p, bold_t + " ", 12, INK, bold=True)
    run(p, rest, 12, BODY)
panel(s, 7.15, 5.55, 5.4, 1.0, fill=LIGHT_AMBER_BG, border=LIGHT_AMBER_BORDER)
_, tf = box(s, 7.3, 5.62, 5.1, 0.9)
p = tf.paragraphs[0]
run(p, "The gap this study works in: ", 11.5, INK, bold=True)
run(p, "every surface above ranks what to buy inside a category the customer already trusts. "
       "None of them argue why to try a category a bad order taught them to avoid.", 11.5, BODY)
_, tf = box(s, 7.15, 6.58, 5.4, 0.52)
run(tf.paragraphs[0], "Discovery surfaces described from the public app and published reporting, "
                      "not internal Blinkit documentation.", 9.5, FOOTNOTE)
notes(s, "The metric: share of monthly active customers buying from at least one new category "
         "that month. Four parts, both artefacts are live and public.")

# ========================================================================
# SLIDE 2 — Why it matters
# ========================================================================
s = slide()
eyebrow(s, "The goal, framed")
h2(s, "Heavy repeat buyers keep narrow baskets — and most of them say so themselves", size=25, y=0.78)
panel(s, 0.55, 2.15, 5.6, 4.6, fill=INK, border=None)
_, tf = box(s, 0.85, 2.4, 5.0, 1.6, anchor=MSO_ANCHOR.TOP)
run(tf.paragraphs[0], "58%", 60, YELLOW, bold=True)
_, tf = box(s, 0.85, 3.75, 5.0, 1.6)
p = tf.paragraphs[0]
run(p, "of surveyed users self-report that they mostly stick to the same categories", 17, WHITE, bold=True)
_, tf = box(s, 0.85, 5.9, 5.0, 0.7)
run(tf.paragraphs[0], "Source: primary survey, 18 of 31 responses. Survey-based; no live interviews.",
    11, DARK_TEXT_MUTED)

panel(s, 6.35, 2.15, 6.45, 2.2, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 6.6, 2.3, 6.0, 1.9)
p = tf.paragraphs[0]
run(p, "This is a trust problem, not an awareness problem", 15, INK, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(6)
run(p2, "These customers order weekly or daily. They have already seen the other categories in "
        "the app. What stops them is what happened — or what they expect to happen — when the "
        "order arrives.", 12.5, BODY)

panel(s, 6.35, 4.5, 6.45, 2.55, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 6.6, 4.62, 6.0, 2.35)
p = tf.paragraphs[0]
run(p, "The company's own growth path runs through this behaviour", 15, INK, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(5)
run(p2, "Analyst, Eternal Q4 FY26 earnings call (28 Apr 2026): “a large part of our growth "
        "narrative from here on depends in some sense on either growing the non-grocery assortment "
        "and going outside of the metro cities.”", 11, BODY)
p3 = tf.add_paragraph(); p3.space_before = Pt(5)
run(p3, "CFO: “It's a function of assortment expansion, geographical expansion as well as more "
        "demand densification in the cities where we are present today.”", 11, BODY)
p4 = tf.add_paragraph(); p4.space_before = Pt(5)
run(p4, "Caveat: ", 10.5, AMBER_TEXT, bold=True)
run(p4, "the call does not discuss product-quality or refund friction. The link to this project's "
        "cause is ours, not the company's.", 10.5, BODY)
notes(s, "Segment: heavy repeat buyers with narrow baskets. 58% self-report sticking to the same "
         "categories. The blocker is trust, not awareness — they have seen the categories, they "
         "distrust them.")

# ========================================================================
# SLIDE 3 — Discovery engine
# ========================================================================
s = slide()
eyebrow(s, "Part 1 · how the engine works", size=12)
_, tf = box(s, 0.55, 0.75, 9.3, 1.1)
run(tf.paragraphs[0], "A five-stage pipeline hard-gates on category-adoption relevance before "
                      "it extracts anything", 21, INK, bold=True)
panel(s, 10.0, 0.8, 2.8, 0.9, fill=LIGHT_AMBER_BG, border=YELLOW)
_, tf = box(s, 10.15, 0.86, 2.5, 0.8)
run(tf.paragraphs[0], "LIVE · DISCOVERY WORKFLOW", 9, AMBER_TEXT, bold=True)
p2 = tf.add_paragraph()
run(p2, "Open the live extractor →", 12, INK, bold=True)

panel(s, 0.55, 1.95, 12.25, 1.0, fill=WHITE, border=CARD_BORDER)
funnel_vals = [("15,820", "public reviews ingested"), ("4,740", "pass the rule-based prefilter"),
               ("1,094", "gate-passing extractions")]
fx = 0.75
for i, (val, label) in enumerate(funnel_vals):
    _, tf = box(s, fx, 2.03, 2.0, 0.85)
    p = tf.paragraphs[0]
    run(p, val, 20, INK if i < 2 else AMBER_TEXT, bold=True)
    p2 = tf.add_paragraph()
    run(p2, label, 9.5, MUTED, bold=True)
    fx += 2.15
    if i < 2:
        _, tfa = box(s, fx - 0.35, 2.2, 0.4, 0.4)
        run(tfa.paragraphs[0], "→", 16, RGBColor(0xB9, 0xA4, 0x5A), bold=True)
_, tf = box(s, 7.3, 2.02, 5.3, 0.75)
p = tf.paragraphs[0]
run(p, "Why gate at all? ", 11.5, INK, bold=True)
run(p, "Hand-reading reviews is anecdote, not evidence — an ungated LLM pass returns a complaints "
       "dump. The gate forces every kept review through one schema.", 11, BODY)

steps = [
    ("1", "Ingest", "15,820 reviews", ["Public app-store reviews", "Deduplicated before any model runs"], WHITE, INK),
    ("2", "Heuristic prefilter", "→ 4,740 kept", ["Drops empty text / under five words", "Drops 5-star reviews", "Rule-based only — no LLM"], WHITE, INK),
    ("3", "Gated extraction", "→ 1,094 kept", ["Gate: trying or avoiding a category?", "Extract: behaviour, category, friction", "Two steps, never one"], INK, WHITE),
    ("4", "Deterministic clustering", "Themes = DB rows", ["Frequency counted in the database", "Never an LLM estimate", "70% land in a single theme"], WHITE, INK),
    ("5", "LLM-judge validation", "Held-out sample", ["Gate decisions re-checked", "Validation dial in the app itself"], WHITE, INK),
]
sx = 0.55
sw = 2.36
for num, name, sub, bullets, fill, tcol in steps:
    panel(s, sx, 2.95, sw - 0.1, 2.05, fill=fill, border=None if fill == INK else CARD_BORDER)
    _, tf = box(s, sx + 0.15, 3.03, sw - 0.35, 1.9)
    p = tf.paragraphs[0]
    run(p, f"{num}  ", 12, YELLOW if fill == INK else AMBER_TEXT, bold=True)
    run(p, name, 12.5, tcol, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(2)
    run(p2, sub, 10, YELLOW if fill == INK else AMBER_TEXT, bold=True)
    for b in bullets:
        pb = tf.add_paragraph(); pb.space_before = Pt(2)
        run(pb, "✓ " + b, 9.5, DARK_TEXT_MUTED2 if fill == INK else BODY)
    sx += sw

# two screenshots
picture_fit(s, IMG / "disc-01-extractor-pass-crop.png", 0.55, 5.05, 3.6, 1.65)
_, tf = box(s, 0.55, 6.75, 3.6, 0.35)
run(tf.paragraphs[0], "✓ KEPT — CATEGORY BEHAVIOUR", 9.5, GREEN_TEXT, bold=True)
picture_fit(s, IMG / "disc-02-extractor-reject-crop.png", 4.3, 5.05, 3.6, 1.65)
_, tf = box(s, 4.3, 6.75, 3.6, 0.35)
run(tf.paragraphs[0], "✗ DROPPED — SERVICE COMPLAINT", 9.5, AMBER_TEXT, bold=True)
panel(s, 8.15, 5.05, 4.65, 0.95, fill=LIGHT_AMBER_BG, border=YELLOW)
_, tf = box(s, 8.3, 5.1, 4.35, 0.85)
p = tf.paragraphs[0]
run(p, "Design decision · the gate. ", 10.5, AMBER_TEXT, bold=True)
run(p, "An early ungated run only surfaced generic service complaints.", 10.5, INK, bold=True)
panel(s, 8.15, 6.1, 4.65, 0.95, fill=LIGHT_GREEN_BG, border=LIGHT_GREEN_BORDER)
_, tf = box(s, 8.3, 6.15, 4.35, 0.85)
p = tf.paragraphs[0]
run(p, "What the engine measures: ", 10.5, GREEN_TEXT, bold=True)
run(p, "adoption behaviour, category named, friction theme. Groq llama-3.1-8b-instant throughout.",
    10.5, BODY)
notes(s, "Required deliverable: the workflow explanation. Five stages. The design decision that "
         "matters is the relevance gate — an early ungated run only surfaced generic service "
         "complaints.")

# ========================================================================
# SLIDE 4 — Review findings
# ========================================================================
s = slide()
eyebrow(s, "Part 1 · what the reviews said")
h2(s, "Poor quality and unreliable products is the dominant reason customers stop widening "
      "their baskets", size=21, y=0.78, w=12.2)

cards4 = [
    ("1", "One theme carries the segment",
     "Of the 1,094 gate-passing extractions, the dominant friction theme is "
     "“Poor Quality and Unreliable Products”.", "70%", "of kept evidence sits in this one theme"),
    ("2", "The steep funnel is the gate working, not data loss",
     "The prefilter drops empty/short/5-star reviews. The gate then removes everything unrelated "
     "to trying or avoiding a category.", "15,820 → 4,740 → 1,094", ""),
]
cx = 0.55
for num, title_t, body_t, stat_v, stat_l in cards4:
    panel(s, cx, 1.85, 4.0, 2.15, fill=WHITE, border=CARD_BORDER)
    _, tf = box(s, cx + 0.2, 1.98, 3.6, 1.95)
    p = tf.paragraphs[0]
    run(p, f"{num}  ", 12, AMBER_TEXT, bold=True)
    run(p, title_t, 13, INK, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(4)
    run(p2, body_t, 10.5, BODY)
    p3 = tf.add_paragraph(); p3.space_before = Pt(6)
    run(p3, stat_v + "  ", 20, GREEN_ACCENT, bold=True)
    run(p3, stat_l, 9.5, MUTED, bold=True)
    cx += 4.15

panel(s, 8.85, 1.85, 3.95, 2.0, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 9.05, 1.98, 3.6, 1.85)
p = tf.paragraphs[0]
run(p, "3  ", 12, AMBER_TEXT, bold=True)
run(p, "Where the top theme lands", 13, INK, bold=True)
cat_bars = [("Groceries", 431, "100%"), ("Electronics", 96, "22%"),
            ("Snacks & beverages", 79, "18%"), ("Household essentials", 39, "9%"),
            ("Personal care", 27, "6%")]
for name, count, pct in cat_bars:
    p = tf.add_paragraph(); p.space_before = Pt(3)
    run(p, f"{name}: ", 9.5, BODY)
    run(p, str(count), 9.5, INK, bold=True)

cards4b = [
    ("4", "Every count is a database row",
     "Frequencies come from the pipeline export — never an LLM estimate, never a per-user "
     "confidence score. The validation dial re-checks gate decisions on a held-out sample."),
]
panel(s, 0.55, 4.0, 4.0, 1.55, fill=LIGHT_GREEN_BG, border=LIGHT_GREEN_BORDER)
_, tf = box(s, 0.75, 4.12, 3.6, 1.35)
p = tf.paragraphs[0]
run(p, "4  ", 12, GREEN_TEXT, bold=True)
run(p, "Every count is a database row", 13, INK, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(4)
run(p2, "Frequencies come from the pipeline export — never an LLM estimate, never a per-user "
        "confidence score.", 10.5, BODY)

picture_fit(s, IMG / "disc-04-results-explorer-crop.png", 8.85, 4.05, 3.95, 2.75)
_, tf = box(s, 8.85, 6.85, 3.95, 0.5)
run(tf.paragraphs[0], "Results Explorer — funnel, ranked themes, source mix, behaviour split and "
                      "the LLM-judge validation dial.", 9, FOOTNOTE)

panel(s, 0.55, 5.75, 7.9, 1.15, fill=INK, border=None)
_, tf = box(s, 0.8, 5.9, 7.4, 0.9, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
run(p, "THE CORE PROBLEM   ", 11, YELLOW, bold=True)
run(p, "Customers are not missing the other categories — they have decided those categories are "
       "unsafe to buy here. One damaged, expired or fake item is read as evidence about the whole "
       "category.", 12.5, WHITE, bold=True)
notes(s, "70% of gate-passing evidence is one theme: poor quality and unreliable products. The "
         "steep funnel is the gate working, not data loss.")

# ========================================================================
# SLIDE 5 — Research: confirmed & challenged
# ========================================================================
s = slide()
eyebrow(s, "Part 2 · primary research")
h2(s, "The survey confirmed the AI's cause — and challenged how far it reaches", size=22, y=0.78, w=8.6)
panel(s, 9.3, 0.55, 3.5, 1.4, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 9.45, 0.62, 3.2, 1.25)
p = tf.paragraphs[0]
run(p, "SAMPLE, STATED HONESTLY", 9.5, AMBER_TEXT, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(3)
run(p2, "31 survey responses. No live interviews — a deliberate trade of depth for breadth. Six "
        "respondents consented to follow-up calls that were not run.", 9.5, INK, bold=True)

# column 1: quantitative validation
panel(s, 0.55, 2.05, 3.95, 4.75, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 0.75, 2.18, 3.6, 4.5)
p = tf.paragraphs[0]
run(p, "%  ", 12, AMBER_TEXT, bold=True)
run(p, "Quantitative validation", 13.5, INK, bold=True)
donuts = [("58%", "Stick to the same categories", "18 of 31"),
          ("48%", "Had a bad order in 3 months", "15 of 31"),
          ("73%", "Of those, fully refunded", "11 of 15")]
for val, label, sub in donuts:
    p = tf.add_paragraph(); p.space_before = Pt(10)
    run(p, val + "  ", 18, GREEN_ACCENT, bold=True)
    run(p, label, 11, INK, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(0)
    run(p2, sub, 9.5, MUTED)
p = tf.add_paragraph(); p.space_before = Pt(12)
run(p, "Asked directly, 1–5: mean 3.0", 10.5, AMBER_TEXT, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(3)
run(p2, "■ 11 agree   ■ 9 neutral   ■ 11 disagree", 10.5, INK, bold=True)
p3 = tf.add_paragraph(); p3.space_before = Pt(3)
run(p3, "Does a bad experience in one category make you hesitant to try others?", 9.5, MUTED)

# column 2: what the survey challenged (dark)
panel(s, 4.65, 2.05, 4.0, 4.75, fill=INK, border=None)
_, tf = box(s, 4.85, 2.18, 3.65, 4.5)
p = tf.paragraphs[0]
run(p, "!  ", 12, YELLOW, bold=True)
run(p, "What the survey challenged", 13.5, YELLOW, bold=True)
challenges = [
    ("The behaviour change is an even split, not a rule.",
     "Of 15 bad orders: 7 changed behaviour, 7 didn't, 1 ambiguous — even split held as the "
     "sample grew."),
    ("Half the “stuck” group had no incident at all.",
     "8 of 18 “stuck” users had no bad order — low intent, not lost trust."),
    ("Some churn instead of narrowing.",
     "“I have started using Zepto more” — platform switching, not category avoidance."),
    ("Packaging is a distinct sub-cause.",
     "A fulfilment failure, not product sourcing — written in unprompted."),
]
for t, d in challenges:
    p = tf.add_paragraph(); p.space_before = Pt(7)
    run(p, t, 11, WHITE, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(2)
    run(p2, d, 10, DARK_TEXT_MUTED2)
p = tf.add_paragraph(); p.space_before = Pt(7)
run(p, "2 respondents raised pricing versus competitors unprompted.", 10, DARK_TEXT_MUTED)

# column 3: confirmed in their words
panel(s, 8.75, 2.05, 4.05, 4.75, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 8.95, 2.18, 3.7, 4.5)
p = tf.paragraphs[0]
run(p, "✓  ", 12, GREEN_TEXT, bold=True)
run(p, "Confirmed, in their words", 13.5, GREEN_TEXT, bold=True)
quotes = [
    ("“I have stopped exploring new categories on Blinkit. I only order the essential items "
     "now if anything urgent.”", "Unresolved puja-item complaint"),
    ("“yes I am skeptical of buying electronic products from Blinkit and other similar "
     "applications”", "Damaged electronics"),
    ("“I avoid buying any fresh products dairy or perishables items from quick commerce "
     "apps.”", "Expired dairy — generalised beyond Blinkit"),
]
for q, src in quotes:
    p = tf.add_paragraph(); p.space_before = Pt(8)
    run(p, q, 10.5, INK, italic=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(2)
    run(p2, src, 9.5, GREEN_TEXT, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(8)
run(p, "The 15 bad orders: 11 damaged · 2 expired · 1 fake/duplicate · 1 other. Resolution: 11 "
       "full refunds · 3 replacements · 1 unresolved.", 9, BODY)

panel(s, 0.55, 6.9, 12.25, 0.55, fill=LIGHT_AMBER_BG, border=YELLOW)
_, tf = box(s, 0.75, 6.93, 11.9, 0.5, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
run(p, "THE VERDICT   ", 10.5, AMBER_TEXT, bold=True)
run(p, "Confirmed: ", 11, INK, bold=True)
run(p, "a quality failure does generalise into avoiding the whole category.  ", 11, INK)
run(p, "Challenged: ", 11, INK, bold=True)
run(p, "it does not do so for everyone — so the fix is scoped to the trust-driven share.", 11, INK)
notes(s, "Mandatory slide: the research both confirmed and challenged the AI. Sample is 31 survey "
         "responses, zero live interviews — a deliberate trade of depth for breadth.")

# ========================================================================
# SLIDE 6 — Problem framing canvas
# ========================================================================
s = slide()
eyebrow(s, "Part 3 · problem framing canvas")
h2(s, "One quality failure generalises into permanent avoidance of an entire category", size=22, y=0.78)

cells = [
    ("01 · WHO IT'S FOR", "Heavy, habitual repeat buyers.",
     "High-frequency customers who default to the same narrow set of categories every month. "
     "18 of 31 surveyed describe themselves this way.", WHITE, AMBER_TEXT),
    ("02 · THE ROOT CAUSE", "Trust, not discovery.",
     "A damaged, expired or fake item teaches the customer that a category is unsafe to buy "
     "here — so they retreat to essentials, or leave for another app.", WHITE, AMBER_TEXT),
    ("03 · HOW THEY COPE TODAY", "Two workarounds, both costly.",
     "Retreat to essentials only, or switch platforms outright. The first suppresses category "
     "adoption silently; the second is active share loss.", WHITE, AMBER_TEXT),
    ("04 · WHY IT'S WORTH IT FOR THE USER", "They want to consolidate — they just can't risk it.",
     "27 of 31 named at least one concrete change that would let them — only 4 said “nothing "
     "in particular” — led by a refund guarantee (13) and a visible quality signal (10).",
     LIGHT_GREEN_BG, GREEN_TEXT),
    ("05 · WHY IT'S WORTH IT FOR THE BUSINESS", "Basket width is the growth lever the company has named.",
     "Restoring trust re-opens non-grocery categories for customers who already order "
     "frequently — no new acquisition required.", LIGHT_GREEN_BG, GREEN_TEXT),
    ("SCOPE LIMIT · STATED UP FRONT", "This addresses only the quality- and trust-driven share of stagnation.",
     "Roughly half the stuck segment reported no incident at all. Those customers are low "
     "intent; no trust fix will move them, and we do not claim it will.", INK, YELLOW),
]
gx, gy = 0.55, 1.75
gw, gh = 4.0, 2.35
for i, (eyebrow_t, title_t, body_t, fill, ecolor) in enumerate(cells):
    col = i % 3
    row = i // 3
    x = gx + col * (gw + 0.15)
    y = gy + row * (gh + 0.15)
    tcol = WHITE if fill == INK else INK
    bcol = DARK_TEXT_MUTED2 if fill == INK else BODY
    panel(s, x, y, gw, gh, fill=fill, border=None if fill != WHITE else CARD_BORDER)
    _, tf = box(s, x + 0.2, y + 0.15, gw - 0.4, gh - 0.3)
    p = tf.paragraphs[0]
    run(p, eyebrow_t, 10, ecolor, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(4)
    run(p2, title_t, 13.5, tcol, bold=True)
    p3 = tf.add_paragraph(); p3.space_before = Pt(4)
    run(p3, body_t, 10.5, bcol)

panel(s, 0.55, 6.65, 12.25, 0.6, fill=LIGHT_AMBER_BG, border=LIGHT_AMBER_BORDER)
_, tf = box(s, 0.75, 6.68, 11.9, 0.55, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
run(p, "Evidence base: ", 10.5, INK, bold=True)
run(p, "1,094 gate-passing review extractions (70% quality-driven) and 31 survey responses "
       "(15 with a recent bad order) point at the same failure.", 10.5, BODY)
notes(s, "Five required framing elements: segment, root cause, workarounds, user value, business "
         "value — plus the scope limit.")

# ========================================================================
# SLIDE 7 — Why trust, not discount
# ========================================================================
s = slide()
eyebrow(s, "Part 4 · solution choice")
h2(s, "A trust-led nudge wins because it answers the blocker customers actually named", size=22, y=0.78, w=12.2)

panel(s, 0.55, 1.85, 8.6, 4.55, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 0.75, 1.98, 8.2, 0.35)
run(tf.paragraphs[0], "OPTIONS CONSIDERED", 12, AMBER_TEXT, bold=True)
headers = ["Option", "Pain it solves", "Whole segment?", "Differentiated?"]
rows_data = [
    ("Discount or intro offer on a new category", "Price hesitation only — 2 respondents raised pricing",
     "✗ No — misses the trust blocker", "✗ No — every rival can copy it"),
    ("Better merchandising and category placement", "Awareness — ranked 3rd at 7% of evidence",
     "✗ No — these buyers already see the categories", "✗ Partly only"),
    ("Let users inspect before accepting delivery", "Pre-purchase doubt — now the 3rd-ranked driver, 7 of 31",
     "✗ Partly — helps at the door, not at the moment of choosing", "✗ No — an ops change any rival can match"),
    ("Trust-led category nudge  ← chosen", "The top two named blockers: refund certainty 13/31, then quality 10/31",
     "✓ Yes for the trust-driven share", "✓ Yes — per-user, per-friction copy"),
]
rows = []
for i, r in enumerate(rows_data):
    bold = i == 3
    color = GREEN_TEXT if bold else INK
    rows.append([(r[0], INK, bold), (r[1], BODY, False), (r[2], color, bold), (r[3], color, bold)])
simple_table(s, 0.75, 2.4, 8.2, 2.9, headers, rows, col_widths=[2.4, 2.6, 2.1, 2.1],
             row_fill={3: LIGHT_AMBER_BG}, body_size=9.5, header_size=10.5)
_, tf = box(s, 0.75, 5.5, 8.2, 0.85)
run(tf.paragraphs[0], "Ranked survey drivers, Q15 (pick up to 2), N=31: no-questions-asked refund "
                      "or return guarantee 13 · visible quality/freshness guarantee 10 · inspect "
                      "before accepting 7 · item reviews/ratings 5 · intro offers 5 · nothing in "
                      "particular 4 · human support agent 3 · safer packaging 1 (write-in).",
    9, FOOTNOTE)

panel(s, 9.35, 1.85, 3.45, 4.55, fill=INK, border=None)
_, tf = box(s, 9.55, 2.0, 3.05, 4.3)
p = tf.paragraphs[0]
run(p, "CHOSEN SOLUTION", 11, YELLOW, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(6)
run(p2, "The Category Nudge Agent", 18, WHITE, bold=True)
p3 = tf.add_paragraph(); p3.space_before = Pt(8)
run(p3, "A per-user nudge into an adjacent category that leads with the refund or return "
        "guarantee, then the quality and freshness signal — in that order, because that is the "
        "order customers ranked them.", 11.5, DARK_TEXT_MUTED2)
p4 = tf.add_paragraph(); p4.space_before = Pt(14)
run(p4, "Why not a discount", 12, WHITE, bold=True)
p5 = tf.add_paragraph(); p5.space_before = Pt(4)
run(p5, "A discount pays a customer to repeat the experience that broke their trust. It does not "
        "change the expected outcome of the order.", 11, DARK_TEXT_MUTED2)
notes(s, "Options considered and why the trust-led nudge wins. Grounded in the ranked survey "
         "drivers: refund guarantee 13, quality/freshness 10.")

# ========================================================================
# SLIDE 8 — The MVP, running
# ========================================================================
s = slide()
eyebrow(s, "Part 4 · the MVP, deployed", size=12)
_, tf = box(s, 0.55, 0.75, 8.6, 1.1)
run(tf.paragraphs[0], "The agent runs today: a shopper-facing nudge, its reasoning, and the push "
                      "queue behind it", 20, INK, bold=True)
_, tf = box(s, 0.55, 1.55, 3.3, 0.35)
p = tf.paragraphs[0]
run(p, "Open the live agent → ", 12, INK, bold=True)
panel(s, 9.35, 0.55, 3.45, 1.15, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 9.5, 0.62, 3.15, 1.0)
p = tf.paragraphs[0]
run(p, "DATA DISCLOSURE", 9.5, AMBER_TEXT, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(3)
run(p2, "8 synthetic user profiles, labelled SYNTHETIC in the product UI itself. No real "
        "customer data was available.", 10, INK, bold=True)

picture_fit(s, IMG / "mvp-02-phone-nudge.png", 0.55, 1.95, 2.35, 4.55)
_, tf = box(s, 0.55, 6.55, 2.35, 0.9)
run(tf.paragraphs[0], "The shopper's view — a trust-led nudge into one adjacent category.", 9, FOOTNOTE)

picture_fit(s, IMG / "mvp-03-reasoning-ranked.png", 3.05, 1.95, 4.85, 4.55)
_, tf = box(s, 3.05, 6.55, 4.85, 0.55)
run(tf.paragraphs[0], "The agent's reasoning panel — why this user, why this category, and what "
                      "else was considered.", 9, FOOTNOTE)

picture_fit(s, IMG / "mvp-07-lockscreen.png", 8.05, 1.95, 2.15, 4.55)
_, tf = box(s, 8.05, 6.55, 2.15, 0.55)
run(tf.paragraphs[0], "Push payloads from the auto-nudge queue.", 9, FOOTNOTE)

annotations = [
    ("Leads with the two research-ranked trust drivers", 10.35, 2.0),
    ("No invented pricing — labelled illustrative", 10.35, 3.1),
    ("Deterministic ranker, integer weights — not model confidence", 10.35, 4.2),
    ("Copy generated live per user, not templated", 10.35, 5.05),
    ("Five users, four different categories — no collapse to one suggestion", 10.35, 5.9),
]
for text, x, y in annotations:
    panel(s, x, y, 2.75, 0.85, fill=WHITE, border=CARD_BORDER)
    _, tf = box(s, x + 0.1, y + 0.05, 2.55, 0.75, anchor=MSO_ANCHOR.MIDDLE)
    run(tf.paragraphs[0], text, 9.5, INK, bold=True)
notes(s, "Annotated teardown of the deployed agent. All profiles are synthetic and labelled as "
         "such in the product UI.")

# ========================================================================
# SLIDE 9 — How it works / breaks
# ========================================================================
s = slide()
eyebrow(s, "Part 4 · architecture and limits")
h2(s, "Deterministic logic decides what to nudge; the LLM only writes the words", size=23, y=0.78, w=12.2)

picture_fit(s, IMG / "mvp-04-out-of-scope.png", 0.55, 1.85, 5.0, 4.35)
panel(s, 0.55, 6.3, 5.0, 0.95, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 0.7, 6.38, 4.7, 0.8)
p = tf.paragraphs[0]
run(p, "The product itself says when the fix does NOT apply. ", 10, INK, bold=True)
run(p, "Amber banner, top of the panel — a low-intent user is told the trust fix isn't theirs.",
    10, BODY)

_, tf = box(s, 5.85, 1.85, 6.9, 0.3)
run(tf.paragraphs[0], "RUNTIME FLOW", 11.5, AMBER_TEXT, bold=True)
flow = [
    ("Input", "User profile — cadence, tenure, category history, prior incident", WHITE, AMBER_TEXT),
    ("Rule-based", "Friction match — profile to a Part 1 theme, no LLM", LIGHT_GREEN_BG, GREEN_TEXT),
    ("Rule-based", "Adjacency ranker — integer weights produce the candidate list", LIGHT_GREEN_BG, GREEN_TEXT),
    ("LLM", "Nudge generation — live Groq call writes copy and reasoning", INK, YELLOW),
    ("Output", "Rendered nudge — refund guarantee first, then quality and freshness", WHITE, AMBER_TEXT),
]
fy = 2.2
for tag, desc, fill, tagcolor in flow:
    tcol = WHITE if fill == INK else INK
    panel(s, 5.85, fy, 6.9, 0.5, fill=fill, border=None if fill != WHITE else CARD_BORDER)
    _, tf = box(s, 6.0, fy + 0.06, 6.6, 0.4, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, tag + "   ", 10.5, tagcolor, bold=True)
    run(p, desc, 11, tcol, bold=True)
    fy += 0.57

panel(s, 5.85, 5.1, 6.9, 1.4, fill=LIGHT_AMBER_BG, border=YELLOW)
_, tf = box(s, 6.0, 5.18, 6.6, 1.25)
p = tf.paragraphs[0]
run(p, "Two integrity rules, enforced in the prompt", 10.5, AMBER_TEXT, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(3)
run(p2, "Never invent statistics. ", 9.5, INK, bold=True)
run(p2, "No ratings, buyer counts or success rates may appear in nudge copy.", 9.5, BODY)
p3 = tf.add_paragraph(); p3.space_before = Pt(2)
run(p3, "Never overclaim to low-intent users. ", 9.5, INK, bold=True)
run(p3, "With no incident on file, the agent states the trust fix does not apply.", 9.5, BODY)

panel(s, 5.85, 6.6, 6.9, 0.85, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 6.0, 6.65, 6.6, 0.75)
p = tf.paragraphs[0]
run(p, "Where it breaks:  ", 10, AMBER_TEXT, bold=True)
run(p, "synthetic profiles, not real history · refund promise inherits fulfilment risk · "
       "packaging failures need an ops fix, not a message · switched-platform users may never "
       "see the nudge.", 9.5, BODY)
notes(s, "Deterministic where correctness matters, LLM only for language. The product itself "
         "states when the fix does not apply.")

# ========================================================================
# SLIDE 10 — Measurement plan
# ========================================================================
s = slide()
eyebrow(s, "Measurement · proposed, not yet measured", size=12)
_, tf = box(s, 0.55, 0.75, 8.4, 1.1)
run(tf.paragraphs[0], "We would judge this on new-category adoption, in three tiers, with trust "
                      "guardrails", 21, INK, bold=True)
panel(s, 9.35, 0.6, 3.45, 1.15, fill=YELLOW, border=None)
_, tf = box(s, 9.5, 0.67, 3.15, 1.0)
run(tf.paragraphs[0], "Plan only. No experiment has been run and the MVP has never been shipped "
                      "to real users, so no result numbers appear on this slide.", 10.5, INK, bold=True)

panel(s, 0.55, 1.95, 4.9, 1.9, fill=INK, border=None)
_, tf = box(s, 0.75, 2.08, 4.5, 1.65)
p = tf.paragraphs[0]
run(p, "PRIMARY METRIC", 10.5, YELLOW, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(4)
run(p2, "% of monthly active customers who purchase from at least one new category that month",
    14.5, WHITE, bold=True)
p3 = tf.add_paragraph(); p3.space_before = Pt(4)
run(p3, "Measured against the nudged segment's own prior month, not against a headline average.",
    10, DARK_TEXT_MUTED)

panel(s, 0.55, 4.0, 4.9, 2.4, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 0.75, 4.13, 4.5, 2.15)
p = tf.paragraphs[0]
run(p, "HONEST SCOPE LIMIT", 10.5, AMBER_TEXT, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(4)
run(p2, "This moves the trust-driven share of stagnation only.", 13.5, INK, bold=True)
p3 = tf.add_paragraph(); p3.space_before = Pt(5)
run(p3, "Around half the stuck segment reported no incident. Their stagnation is low intent, and "
        "a trust message is the wrong instrument for it. Any evaluation must report the two "
        "sub-segments separately or it will read as a failure of the wrong thing.", 11, BODY)

panel(s, 5.6, 1.95, 7.2, 2.3, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 5.8, 2.05, 6.8, 0.3)
run(tf.paragraphs[0], "MEASUREMENT LADDER", 11, AMBER_TEXT, bold=True)
headers10 = ["Tier", "What we would watch", "Reads as"]
rows10 = [
    [("1 · Outcome", INK, True), ("New-category purchase rate among nudged trust-blocked users", BODY, False), ("Did the behaviour change", INK, True)],
    [("2 · Mechanism", INK, True), ("Nudge shown → opened → category browsed → item added", BODY, False), ("Which step leaks", INK, True)],
    [("3 · Durability", INK, True), ("Repeat purchase in the newly tried category the following month", BODY, False), ("Trust restored, not just tested", INK, True)],
]
simple_table(s, 5.8, 2.4, 6.8, 1.75, headers10, rows10, col_widths=[1.4, 2.9, 1.9], body_size=9.5, header_size=10.5)

panel(s, 5.6, 4.4, 7.2, 2.0, fill=WHITE, border=CARD_BORDER)
_, tf = box(s, 5.8, 4.5, 6.8, 1.85)
p = tf.paragraphs[0]
run(p, "GUARDRAILS", 11, AMBER_TEXT, bold=True)
guardrails = [
    ("Return and refund rate", "in nudged categories must not rise — a nudge into a genuinely bad category is a trust loss, not a win."),
    ("Complaint rate", "among nudged users, tracked against their own baseline."),
    ("Order frequency", "and nudge fatigue — reorder rate must not fall; opt-outs capped per user."),
    ("Copy audit", "sampled nudges checked for invented statistics; any hit is a release blocker."),
]
for i in range(0, 4, 2):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    if i == 0:
        p = tf.add_paragraph()
    p.space_before = Pt(6)
    b, d = guardrails[i]
    run(p, b + " ", 10, INK, bold=True)
    run(p, d, 10, BODY)
    b2, d2 = guardrails[i + 1]
    p2 = tf.add_paragraph(); p2.space_before = Pt(4)
    run(p2, b2 + " ", 10, INK, bold=True)
    run(p2, d2, 10, BODY)

_, tf = box(s, 0.55, 6.5, 5.0, 0.3)
run(tf.paragraphs[0], "WHAT WE'D BUILD NEXT", 11, AMBER_TEXT, bold=True)
next_items = [
    ("1 · Let users inspect before accepting", "The 3rd-ranked driver (7 of 31) — an operations change, not a message."),
    ("2 · A relevance-led nudge for the low-intent half", "8 of 18 stuck users had no bad experience at all."),
    ("3 · Real behavioural data instead of synthetic profiles", "The first real test is whether the friction match holds."),
]
nx = 0.55
for title_t, body_t in next_items:
    panel(s, nx, 6.85, 4.0, 0.6, fill=WHITE, border=CARD_BORDER)
    _, tf = box(s, nx + 0.1, 6.87, 3.8, 0.56)
    p = tf.paragraphs[0]
    run(p, title_t, 9, INK, bold=True)
    nx += 4.15
notes(s, "This is a plan. No experiment has been run, so there are no results on this slide — say "
         "that out loud.")

out = HERE / "NL_Blinkit_CategoryAdoption.pptx"
prs.save(out)
print(f"Saved {out} ({len(prs.slides._sldIdLst)} slides)")
